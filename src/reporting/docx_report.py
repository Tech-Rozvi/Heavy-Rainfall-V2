"""
Word (.docx) report writers for ROZVI site hazard and event-validation reports.

Requires python-docx. Falls back with a clear error if missing.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional, Sequence

from src.utils import get_logger

logger = get_logger(__name__)


def _require_docx():
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        return Document, Inches, Pt, RGBColor, WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise ImportError(
            "python-docx is required for Word reports. Install with: pip install python-docx"
        ) from exc


def _set_run_font(run, *, bold=False, size=11, color=None):
    run.bold = bold
    run.font.size = size
    run.font.name = "Arial"
    if color is not None:
        run.font.color.rgb = color


def _add_heading(doc, text: str, level: int = 1):
    Document, Inches, Pt, RGBColor, _ = _require_docx()
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Arial"
        if level == 1:
            run.font.color.rgb = RGBColor(0x1A, 0x56, 0xA0)
    return h


def _add_para(doc, text: str, *, bold=False, italic=False, size=11):
    Document, Inches, Pt, RGBColor, _ = _require_docx()
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Arial"
    return p


def _add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[Any]]):
    Document, Inches, Pt, RGBColor, _ = _require_docx()
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = str(h)
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(10)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = "" if val is None else str(val)
            for p in cells[c_i].paragraphs:
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
    return table


def write_site_hazard_docx(
    *,
    site_name: str,
    bounds: tuple[float, float, float, float],
    mode: str,
    depth_rows: list[dict],
    climate_rows: list[dict],
    exposure: dict,
    exposure_rows: list[dict],
    peril: dict,
    sources: Optional[dict] = None,
    generated_at: Optional[str] = None,
    output_path: str | Path = "outputs/reports/site_hazard_report.docx",
    extra_notes: Optional[list[str]] = None,
) -> Path:
    """
    Write an executive site flood hazard report as Word (.docx).

    Physical hazard only — no financial loss figures.
    """
    Document, Inches, Pt, RGBColor, WD_ALIGN_PARAGRAPH = _require_docx()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    title = doc.add_heading("Site Flood Hazard Report", level=0)
    for run in title.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xA0)

    gen = generated_at or datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
    b = bounds
    bounds_str = f"({b[0]:.4f}, {b[1]:.4f}) – ({b[2]:.4f}, {b[3]:.4f})"
    _add_para(doc, f"Site: {site_name}", bold=True)
    _add_para(doc, f"Generated: {gen} UTC", size=10)
    _add_para(doc, f"AOI bounds: {bounds_str}", size=10)
    _add_para(doc, f"Mode: {mode}", size=10)

    note = doc.add_paragraph()
    run = note.add_run(
        "Physical hazard assessment only. Depths in metres. Building counts are "
        "structure footprints intersecting inundation — not insured loss or monetary damage. "
        "Depth products may be provisional unless explicitly event-validated."
    )
    run.italic = True
    run.font.size = Pt(10)
    run.font.name = "Arial"

    _add_heading(doc, "1. Multi-peril indicators", level=1)
    _add_table(
        doc,
        ["Peril", "Indicator"],
        [
            ["Pluvial", peril.get("pluvial", "n/a")],
            ["Fluvial", peril.get("fluvial", "n/a")],
            ["Coastal", peril.get("coastal", "n/a")],
        ],
    )
    if peril.get("notes"):
        _add_para(doc, str(peril.get("notes")), size=10)

    _add_heading(doc, "2. Peak water depth by return period", level=1)
    if depth_rows:
        _add_table(
            doc,
            ["Return period T (years)", "Peak depth h_max (m)", "Mean depth where wet (m)"],
            [[r.get("T"), r.get("h_max"), r.get("h_mean_wet")] for r in depth_rows],
        )
    else:
        _add_para(doc, "No depth-by-return-period layers were supplied for this run.", size=10)

    _add_heading(doc, "3. Climate delta scales (transparent factors)", level=1)
    if climate_rows:
        _add_table(
            doc,
            ["SSP", "2030", "2050", "2080"],
            [
                [r.get("ssp"), r.get("y2030"), r.get("y2050"), r.get("y2080")]
                for r in climate_rows
            ],
        )
        _add_para(
            doc,
            "These are transparent delta-change multipliers on design rainfall, "
            "not dynamical downscaling outputs.",
            size=10,
            italic=True,
        )
    else:
        _add_para(doc, "No climate scales provided.", size=10)

    _add_heading(doc, "4. Building exposure (physical counts)", level=1)
    n_b = (exposure or {}).get("n_buildings", 0)
    _add_para(doc, f"Buildings assessed: {n_b}", size=11)
    if exposure_rows:
        _add_table(
            doc,
            ["Depth threshold", "Structures intersecting inundation"],
            [[r.get("threshold"), r.get("count")] for r in exposure_rows],
        )
    if (exposure or {}).get("error"):
        _add_para(doc, f"Exposure note: {exposure.get('error')}", size=10)

    _add_heading(doc, "5. Data sources", level=1)
    src = sources or {}
    if src:
        import json
        _add_para(doc, json.dumps(src, indent=2, default=str)[:4000], size=9)
    else:
        _add_para(doc, "No source metadata attached.", size=10)

    if extra_notes:
        _add_heading(doc, "6. Notes", level=1)
        for n in extra_notes:
            _add_para(doc, f"• {n}", size=10)

    _add_para(
        doc,
        "ROZVI — rainfall / extremes ≠ flood depth ≠ financial risk. "
        "See docs/SCIENTIFIC_CLAIMS.md.",
        size=9,
        italic=True,
    )

    doc.save(str(output_path))
    logger.info("Wrote Word report %s", output_path)
    return output_path


def write_validation_report_docx(
    report: Any,
    output_path: str | Path = "outputs/validation/validation_report.docx",
) -> Path:
    """Write an event depth-validation report to Word."""
    Document, Inches, Pt, RGBColor, WD_ALIGN_PARAGRAPH = _require_docx()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    data = report.to_dict() if hasattr(report, "to_dict") else dict(report)
    doc = Document()
    title = doc.add_heading("Event Depth Validation Report", level=0)
    for run in title.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xA0)

    _add_para(doc, f"Event ID: {data.get('event_id', '')}", bold=True)
    _add_para(doc, f"Validation status: {data.get('validation_status', '')}", bold=True)
    _add_para(doc, f"Passed: {data.get('passed', False)}")
    _add_para(doc, f"High-water marks: {data.get('n_marks', 0)}")

    metrics = data.get("point_metrics") or {}
    _add_heading(doc, "Point metrics", level=1)
    _add_table(
        doc,
        ["Metric", "Value"],
        [[k, metrics.get(k)] for k in ("n_compared", "bias_m", "mae_m", "rmse_m", "max_abs_error_m")],
    )

    comps = data.get("point_comparisons") or []
    if comps:
        _add_heading(doc, "Mark comparisons", level=1)
        _add_table(
            doc,
            ["ID", "x", "y", "obs_m", "sim_m", "error_m", "inside"],
            [
                [
                    c.get("mark_id"),
                    c.get("x"),
                    c.get("y"),
                    c.get("observed_m"),
                    c.get("simulated_m"),
                    c.get("error_m"),
                    c.get("inside_grid"),
                ]
                for c in comps[:50]
            ],
        )

    fp = data.get("footprint")
    if fp:
        _add_heading(doc, "Footprint scores", level=1)
        _add_table(doc, ["Metric", "Value"], [[k, v] for k, v in fp.items()])

    notes = data.get("notes") or []
    if notes:
        _add_heading(doc, "Notes", level=1)
        for n in notes:
            _add_para(doc, f"• {n}", size=10)

    _add_para(
        doc,
        "event_validated applies only to this event depth field when passed=true. "
        "Not regulatory design approval.",
        size=9,
        italic=True,
    )
    doc.save(str(output_path))
    logger.info("Wrote validation Word report %s", output_path)
    return output_path
